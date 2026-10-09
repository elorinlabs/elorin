from pathlib import Path
import zipfile,io,shutil
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from PIL import Image
from docx import Document
from docx.shared import Inches
from pypdf import PdfReader,PdfWriter
from pypdf.annotations import Link,Text
from pypdf.generic import DictionaryObject,NameObject,TextStringObject
from pypdf.constants import UserAccessPermissions

root=Path(__file__).parent/'fixtures'/'documents';root.mkdir(parents=True,exist_ok=True)
image=root/'picture.png';Image.new('RGB',(640,400),'#638cc3').save(image)
def pdf(name,pages=3,scan=False,unicode=False):
 c=canvas.Canvas(str(root/name),pagesize=(612,792));c.setTitle('Prism document fixtures');c.setAuthor('Prism QA')
 if unicode:pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
 for n in range(1,pages+1):
  if scan:c.drawImage(str(image),60,220,width=490,height=300)
  else:
   c.setFont('STSong-Light' if unicode else 'Helvetica',20);c.drawString(60,710,'中文文档 日本語 テスト' if unicode else f'Prism PDF page {n}');c.setFont('Helvetica',12);c.drawString(60,665,f'Needle search result {n}. Selectable text and links.')
  if n<=3:c.bookmarkPage(f'page{n}');c.addOutlineEntry(f'Section {n}',f'page{n}',0)
  if n==1 and not scan:c.linkURL('https://example.com', (60,610,220,640),relative=0);c.drawString(60,620,'Safe external link');c.linkRect('', 'page1',(60,580,220,605))
  c.showPage()
 c.save()
pdf('basic.pdf');pdf('long.pdf',300);pdf('outline.pdf',10);pdf('image-only.pdf',2,True);pdf('unicode.pdf',2,unicode=True);shutil.copy(root/'basic.pdf',root/'links.pdf')
w=PdfWriter();w.append(root/'basic.pdf');w.encrypt('prism-secret');w.write(root/'password.pdf')
w=PdfWriter();w.append(root/'basic.pdf');w.encrypt(user_password='',owner_password='owner-only',permissions_flag=UserAccessPermissions.PRINT);w.write(root/'restricted.pdf')
w=PdfWriter();w.append(root/'basic.pdf');w.add_annotation(0,Text(rect=(300,600,340,640),text='Read-only annotation note'));w.add_attachment('notes.txt',b'Never execute attachments');w.add_js("globalThis.PRISM_PDF_EXECUTED=true");w.write(root/'actions.pdf')
(root/'corrupted.pdf').write_bytes(b'%PDF-1.7\nnot a valid document')
def sparse_pdf(name,pages,padding):
 with (root/name).open('wb') as f:
  f.write(b'%PDF-1.7\n');offsets=[0]
  def obj(n,body):
   assert n==len(offsets);offsets.append(f.tell());f.write(f'{n} 0 obj\n'.encode()+body+b'\nendobj\n')
  obj(1,b'<< /Type /Catalog /Pages 2 0 R >>');kids=' '.join(f'{5+n} 0 R' for n in range(pages));obj(2,f'<< /Type /Pages /Count {pages} /Kids [{kids}] >>'.encode());obj(3,b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>');text=b'BT /F1 18 Tf 60 710 Td (Prism large PDF Needle) Tj ET';obj(4,f'<< /Length {len(text)} >>\nstream\n'.encode()+text+b'\nendstream')
  for n in range(pages):obj(5+n,b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> /Contents 4 0 R >>')
  n=len(offsets);offsets.append(f.tell());f.write(f'{n} 0 obj\n<< /Length {padding} >>\nstream\n'.encode());f.seek(padding,1);f.write(b'\nendstream\nendobj\n');xref=f.tell();f.write(f'xref\n0 {len(offsets)}\n0000000000 65535 f \n'.encode());
  for offset in offsets[1:]:f.write(f'{offset:010} 00000 n \n'.encode())
  f.write(f'trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF'.encode())
sparse_pdf('generated-1001.pdf',1001,0);sparse_pdf('generated-100mb.pdf',10,101*1024*1024);sparse_pdf('generated-500mb.pdf',1001,501*1024*1024)
doc=Document();doc.core_properties.title='Prism Office Fixture';doc.core_properties.author='Prism QA';doc.add_heading('Document title',0);doc.add_heading('First heading',1);p=doc.add_paragraph('中文 日本語 Needle searchable paragraph. ');p.add_run('Bold').bold=True;p.add_run(' Italic').italic=True;p.add_run(' Underline').underline=True;doc.add_paragraph('List item',style='List Bullet');table=doc.add_table(rows=3,cols=3)
for i,row in enumerate(table.rows):
 for j,cell in enumerate(row.cells):cell.text=f'Cell {i+1},{j+1} Needle'
doc.add_picture(str(image),width=Inches(3));doc.add_page_break();doc.add_heading('Second heading',2);doc.add_paragraph('Another section.');doc.save(root/'basic.docx')
for name in ['headings.docx','tables.docx','images.docx','links.docx']:shutil.copy(root/'basic.docx',root/name)
with zipfile.ZipFile(root/'basic.docx') as z:
 entries={name:z.read(name) for name in z.namelist()}
text=entries['word/document.xml'].decode();text=text.replace('</w:body>','<w:p><w:hyperlink r:id="rIdPrism"><w:r><w:t>External link</w:t></w:r></w:hyperlink></w:p></w:body>');entries['word/document.xml']=text.encode();relations=entries['word/_rels/document.xml.rels'].decode().replace('</Relationships>','<Relationship Id="rIdPrism" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://example.com" TargetMode="External"/></Relationships>');entries['word/_rels/document.xml.rels']=relations.encode()
with zipfile.ZipFile(root/'links.docx','w',zipfile.ZIP_DEFLATED) as z:
 for name,data in entries.items():z.writestr(name,data)
with zipfile.ZipFile(root/'malformed.docx','w') as z:z.writestr('[Content_Types].xml','<Types/>');z.writestr('word/document.xml','<!DOCTYPE evil SYSTEM "file:///etc/passwd"><broken>')
odt='''<?xml version="1.0"?><office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"><office:body><office:text><text:h text:outline-level="1">ODT heading</text:h><text:p>中文 Needle ODT paragraph.</text:p><text:list><text:list-item><text:p>List item</text:p></text:list-item></text:list><table:table><table:table-row><table:table-cell><text:p>Cell A</text:p></table:table-cell><table:table-cell><text:p>Cell B</text:p></table:table-cell></table:table-row></table:table></office:text></office:body></office:document-content>'''
with zipfile.ZipFile(root/'basic.odt','w') as z:z.writestr('mimetype','application/vnd.oasis.opendocument.text');z.writestr('content.xml',odt)
(root/'basic.rtf').write_text(r'{\rtf1\ansi\uc1 Basic RTF Needle\par {\b Bold} {\i Italic} {\ul Underline}\par Chinese: \u20013?\u25991?}',encoding='ascii')
(root/'legacy.doc').write_bytes(bytes.fromhex('d0cf11e0a1b11ae1')+bytes(504)+'WordDocument'.encode('utf-16le')+bytes(100))
print('Generated Module 09 fixtures:',len(list(root.iterdir())))

# Every JPEG stream below is referenced by a visible page, unlike the sparse range fixture.
def image_heavy_pdf():
 destination=root/'generated-content-500mb.pdf'
 if destination.exists():return
 jpg=io.BytesIO();Image.effect_noise((1400,1400),100).convert('RGB').save(jpg,format='JPEG',quality=96);data=jpg.getvalue();pages=(501*1024*1024)//len(data)+1
 with destination.open('wb') as f:
  f.write(b'%PDF-1.7\n');offsets=[0]
  def obj(n,body):
   assert n==len(offsets);offsets.append(f.tell());f.write(f'{n} 0 obj\n'.encode()+body+b'\nendobj\n')
  obj(1,b'<< /Type /Catalog /Pages 2 0 R >>');kids=' '.join(f'{5+n*2} 0 R' for n in range(pages));obj(2,f'<< /Type /Pages /Count {pages} /Kids [{kids}] >>'.encode());obj(3,b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>');text=b'BT /F1 18 Tf 60 740 Td (Prism image-heavy Needle PDF) Tj ET q 490 0 0 600 60 100 cm /Im1 Do Q';obj(4,f'<< /Length {len(text)} >>\nstream\n'.encode()+text+b'\nendstream')
  for page in range(pages):
   n=5+page*2;obj(n,f'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> /XObject << /Im1 {n+1} 0 R >> >> /Contents 4 0 R >>'.encode());obj(n+1,f'<< /Type /XObject /Subtype /Image /Width 1400 /Height 1400 /ColorSpace /DeviceRGB /BitsPerComponent 8 /Filter /DCTDecode /Length {len(data)} >>\nstream\n'.encode()+data+b'\nendstream')
  xref=f.tell();f.write(f'xref\n0 {len(offsets)}\n0000000000 65535 f \n'.encode())
  for offset in offsets[1:]:f.write(f'{offset:010} 00000 n \n'.encode())
  f.write(f'trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF'.encode())
 print('Image-heavy PDF:',pages,'pages,',destination.stat().st_size,'bytes')
image_heavy_pdf()
longdoc=Document();longdoc.add_heading('Long Office reading fixture',1)
for i in range(3000):
 if i%30==0:longdoc.add_heading(f'Chapter {i//30+1}',2)
 longdoc.add_paragraph(f'Paragraph {i+1}: 中文 Needle text for progressive reading, outline and search. '+ ('ClosingUniqueNeedle' if i==2999 else ''))
longdoc.save(root/'generated-long.docx')
many=Document();many.add_heading('Many images and tables',1)
for i in range(40):
 buffer=io.BytesIO();Image.new('RGB',(800,500),(i*6,80,150)).save(buffer,format='PNG');buffer.seek(0);many.add_picture(buffer,width=Inches(2));table=many.add_table(rows=3,cols=3)
 for row in table.rows:
  for cell in row.cells:cell.text=f'Table {i+1} 中文'
many.save(root/'generated-many.docx')
